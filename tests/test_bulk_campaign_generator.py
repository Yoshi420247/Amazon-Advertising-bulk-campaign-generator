"""Tests for bulk_campaign_generator.py"""

import json
import os
import tempfile
from pathlib import Path

import pandas as pd
import pytest

# Ensure the project root is importable
import sys
sys.path.insert(0, str(Path(__file__).parent.parent))

from bulk_campaign_generator import (
    BulkCampaignGenerator,
    CampaignConfig,
    DEFAULT_COLUMNS,
    ValidationResult,
    _find_column,
    build_asin_sku_mapping,
    generate_bulk_upload,
    get_listing_info,
    read_active_listings_report,
    read_template_headers,
    validate_all_rows,
    validate_bidding_adjustment_row,
    validate_campaign_row,
    validate_keyword_row,
    validate_product_ad_row,
    validate_product_targeting_row,
    write_bulk_xlsx,
)

# =========================================================================
# Fixtures
# =========================================================================

SAMPLES_DIR = Path(__file__).parent.parent / "samples"


@pytest.fixture
def sample_listings_path():
    return str(SAMPLES_DIR / "sample_active_listings.txt")


@pytest.fixture
def sample_asins_path():
    return str(SAMPLES_DIR / "sample_asins.txt")


@pytest.fixture
def sample_config_path():
    return str(SAMPLES_DIR / "sample_config.json")


@pytest.fixture
def sample_tier_path():
    return str(SAMPLES_DIR / "sample_tier_assignment.json")


@pytest.fixture
def listings_df(sample_listings_path):
    return read_active_listings_report(sample_listings_path)


@pytest.fixture
def config():
    return CampaignConfig()


@pytest.fixture
def generator(config):
    return BulkCampaignGenerator(config, DEFAULT_COLUMNS)


# =========================================================================
# _find_column tests
# =========================================================================

class TestFindColumn:
    def test_finds_first_matching_column(self):
        df = pd.DataFrame(columns=["foo", "asin1", "ASIN"])
        assert _find_column(df, ["asin1", "ASIN"], "ASIN") == "asin1"

    def test_finds_fallback_column(self):
        df = pd.DataFrame(columns=["foo", "ASIN"])
        assert _find_column(df, ["asin1", "ASIN"], "ASIN") == "ASIN"

    def test_raises_when_not_found(self):
        df = pd.DataFrame(columns=["foo", "bar"])
        with pytest.raises(ValueError, match="Could not find ASIN column"):
            _find_column(df, ["asin1", "ASIN"], "ASIN")


# =========================================================================
# Active Listings Report tests
# =========================================================================

class TestActiveListingsReport:
    def test_read_tab_delimited(self, sample_listings_path):
        df = read_active_listings_report(sample_listings_path)
        assert len(df) == 7
        assert "asin1" in df.columns
        assert "seller-sku" in df.columns

    def test_read_csv_fallback(self):
        """Test that CSV files are also supported."""
        with tempfile.NamedTemporaryFile(mode="w", suffix=".csv", delete=False) as f:
            f.write("seller-sku,asin1,item-name,price\n")
            f.write("SKU-001,B08NV6CLGF,Test Product,9.99\n")
            f.name
        try:
            df = read_active_listings_report(f.name)
            assert len(df) == 1
            assert df.iloc[0]["asin1"] == "B08NV6CLGF"
        finally:
            os.unlink(f.name)

    def test_file_not_found(self):
        with pytest.raises(FileNotFoundError):
            read_active_listings_report("/nonexistent/file.txt")

    def test_build_asin_sku_mapping(self, listings_df):
        mapping = build_asin_sku_mapping(listings_df)
        assert len(mapping) == 7
        assert mapping["B08NV6CLGF"] == "m_black/w_3x4.5_100"
        assert mapping["B09PQR1234"] == "m_silver/w_10x14_100"

    def test_get_listing_info(self, listings_df):
        info = get_listing_info(listings_df, "B08NV6CLGF")
        assert "item-name" in info
        assert "3x4.5" in info["item-name"]
        assert info["price"] == "12.99"

    def test_get_listing_info_not_found(self, listings_df):
        info = get_listing_info(listings_df, "NONEXISTENT")
        assert info == {}


# =========================================================================
# Template Reader tests
# =========================================================================

class TestTemplateReader:
    def test_default_columns_returned_when_no_template(self):
        headers = read_template_headers(None)
        assert headers == DEFAULT_COLUMNS
        # Ensure it's a copy, not the original
        headers.append("extra")
        assert len(DEFAULT_COLUMNS) != len(headers)

    def test_missing_template_falls_back(self):
        headers = read_template_headers("/nonexistent/template.xlsx")
        assert headers == DEFAULT_COLUMNS


# =========================================================================
# Validation tests
# =========================================================================

class TestValidation:
    def test_valid_campaign_row(self):
        row = {
            "Campaign ID": -1000,
            "Campaign Name": "Test Campaign",
            "Campaign Daily Budget": 50.0,
            "Campaign Targeting Type": "MANUAL",
            "Campaign Status": "enabled",
            "Campaign Start Date": "20250101",
            "Campaign Bidding Strategy": "Dynamic bids - down only",
        }
        result = validate_campaign_row(row)
        assert result.is_valid

    def test_campaign_missing_name(self):
        row = {
            "Campaign ID": -1000,
            "Campaign Name": "",
            "Campaign Daily Budget": 50.0,
            "Campaign Targeting Type": "MANUAL",
            "Campaign Status": "enabled",
            "Campaign Start Date": "20250101",
            "Campaign Bidding Strategy": "Dynamic bids - down only",
        }
        result = validate_campaign_row(row)
        assert not result.is_valid
        assert any("Campaign Name" in e for e in result.errors)

    def test_campaign_invalid_targeting_type(self):
        row = {
            "Campaign ID": -1000,
            "Campaign Name": "Test",
            "Campaign Daily Budget": 50.0,
            "Campaign Targeting Type": "INVALID",
            "Campaign Status": "enabled",
            "Campaign Start Date": "20250101",
            "Campaign Bidding Strategy": "Dynamic bids - down only",
        }
        result = validate_campaign_row(row)
        assert not result.is_valid
        assert any("Targeting Type" in e for e in result.errors)

    def test_campaign_negative_budget(self):
        row = {
            "Campaign ID": -1000,
            "Campaign Name": "Test",
            "Campaign Daily Budget": -10.0,
            "Campaign Targeting Type": "MANUAL",
            "Campaign Status": "enabled",
            "Campaign Start Date": "20250101",
            "Campaign Bidding Strategy": "Dynamic bids - down only",
        }
        result = validate_campaign_row(row)
        assert not result.is_valid
        assert any("positive" in e for e in result.errors)

    def test_valid_bidding_adjustment(self):
        row = {"Campaign ID": -1000, "Placement": "placementTop", "Percentage": 80}
        result = validate_bidding_adjustment_row(row)
        assert result.is_valid

    def test_bidding_adjustment_negative_percentage(self):
        row = {"Campaign ID": -1000, "Placement": "placementTop", "Percentage": -5}
        result = validate_bidding_adjustment_row(row)
        assert not result.is_valid
        assert any("non-negative" in e for e in result.errors)

    def test_bidding_adjustment_exceeds_max(self):
        row = {"Campaign ID": -1000, "Placement": "placementTop", "Percentage": 999}
        result = validate_bidding_adjustment_row(row)
        assert not result.is_valid
        assert any("900%" in e for e in result.errors)

    def test_valid_keyword_row(self):
        row = {
            "Campaign ID": -1000,
            "Ad Group ID": -2000,
            "Keyword Text": "mylar bags",
            "Match Type": "exact",
            "Keyword Status": "enabled",
        }
        result = validate_keyword_row(row)
        assert result.is_valid

    def test_keyword_invalid_match_type(self):
        row = {
            "Campaign ID": -1000,
            "Ad Group ID": -2000,
            "Keyword Text": "mylar bags",
            "Match Type": "INVALID",
            "Keyword Status": "enabled",
        }
        result = validate_keyword_row(row)
        assert not result.is_valid

    def test_valid_product_ad_row(self):
        row = {
            "Campaign ID": -1000,
            "Ad Group ID": -2000,
            "SKU": "test-sku-123",
            "Ad Status": "enabled",
        }
        result = validate_product_ad_row(row)
        assert result.is_valid

    def test_valid_product_targeting_row(self):
        row = {
            "Campaign ID": -1000,
            "Ad Group ID": -2000,
            "Product Targeting Expression": 'asin="B08NV6CLGF"',
            "Product Targeting Status": "enabled",
        }
        result = validate_product_targeting_row(row)
        assert result.is_valid

    def test_validate_all_rows_catches_errors(self):
        rows = [
            {"Entity": "Campaign", "Campaign ID": "", "Campaign Name": ""},
            {"Entity": "Keyword", "Campaign ID": -1000, "Ad Group ID": -2000,
             "Keyword Text": "test", "Match Type": "exact", "Keyword Status": "enabled"},
        ]
        result = validate_all_rows(rows)
        assert not result.is_valid
        # Campaign row should fail, keyword row should pass
        assert len(result.errors) > 0


# =========================================================================
# BulkCampaignGenerator tests
# =========================================================================

class TestBulkCampaignGenerator:
    def test_campaign_ids_start_at_expected_value(self, generator):
        cid = generator.add_campaign("Test", 50.0, "MANUAL")
        assert cid == -1000

    def test_ad_group_ids_start_at_expected_value(self, generator):
        agid = generator.add_ad_group(-1000, "Test AG", 0.75)
        assert agid == -2000

    def test_ids_are_sequential(self, generator):
        cid1 = generator.add_campaign("Test 1", 50.0, "MANUAL")
        cid2 = generator.add_campaign("Test 2", 30.0, "AUTO")
        assert cid1 == -1000
        assert cid2 == -1001

    def test_add_campaign_creates_correct_row(self, generator):
        cid = generator.add_campaign("Test Campaign", 50.0, "MANUAL")
        rows = generator.get_rows()
        assert len(rows) == 1
        row = rows[0]
        assert row["Entity"] == "Campaign"
        assert row["Operation"] == "Create"
        assert row["Campaign Name"] == "Test Campaign"
        assert row["Campaign Daily Budget"] == 50.0
        assert row["Campaign Targeting Type"] == "MANUAL"
        assert row["Campaign Status"] == "enabled"

    def test_add_bidding_adjustment(self, generator):
        generator.add_bidding_adjustment(-1000, "placementTop", 80)
        rows = generator.get_rows()
        assert len(rows) == 1
        row = rows[0]
        assert row["Entity"] == "Bidding Adjustment"
        assert row["Placement"] == "placementTop"
        assert row["Percentage"] == 80

    def test_add_product_ad(self, generator):
        generator.add_product_ad(-1000, -2000, "test-sku")
        rows = generator.get_rows()
        assert len(rows) == 1
        assert rows[0]["SKU"] == "test-sku"
        assert rows[0]["Entity"] == "Product Ad"

    def test_add_keyword(self, generator):
        generator.add_keyword(-1000, -2000, "mylar bags", "exact", bid=1.50)
        rows = generator.get_rows()
        assert len(rows) == 1
        assert rows[0]["Keyword Text"] == "mylar bags"
        assert rows[0]["Match Type"] == "exact"
        assert rows[0]["Bid"] == 1.50

    def test_add_keyword_without_bid(self, generator):
        generator.add_keyword(-1000, -2000, "mylar bags", "exact")
        rows = generator.get_rows()
        assert rows[0]["Bid"] == ""

    def test_add_product_targeting_auto(self, generator):
        generator.add_product_targeting(-1000, -2000, "close-match",
                                        expression_type="auto")
        rows = generator.get_rows()
        assert len(rows) == 1
        assert rows[0]["Product Targeting Expression"] == "close-match"
        assert rows[0]["Product Targeting Expression Type"] == "auto"

    def test_add_product_targeting_manual(self, generator):
        generator.add_product_targeting(-1000, -2000, 'asin="B08NV6CLGF"')
        rows = generator.get_rows()
        assert rows[0]["Product Targeting Expression Type"] == "manual"

    def test_empty_row_has_all_headers(self, generator):
        row = generator._create_empty_row()
        assert set(row.keys()) == set(DEFAULT_COLUMNS)
        assert all(v == "" for v in row.values())

    def test_generate_clearance_campaigns_structure(self, generator):
        """Test that the full campaign generation creates the right structure."""
        skus = ["sku-001", "sku-002"]
        competitor_asins = ["B07COMP001"]

        generator.generate_clearance_campaigns(
            "365+", skus,
            competitor_asins=competitor_asins
        )

        rows = generator.get_rows()

        # Count entities
        entity_counts = {}
        for row in rows:
            entity = row["Entity"]
            entity_counts[entity] = entity_counts.get(entity, 0) + 1

        # Should have 3 campaigns (Manual KW, Auto, Product Targeting)
        assert entity_counts["Campaign"] == 3

        # Each campaign gets 2 bidding adjustments (top + product page)
        assert entity_counts["Bidding Adjustment"] == 6

        # Manual KW: 2 ad groups (one per SKU)
        # Auto: 1 ad group
        # Product Targeting: 1 ad group
        assert entity_counts["Ad Group"] == 4

        # Product ads: Manual KW has 2 (one per SKU) + Auto has 2 + PT has 2
        assert entity_counts["Product Ad"] == 6

        # Auto targeting: 4 groups (close-match, loose-match, substitutes, complements)
        # Product targeting: 1 competitor ASIN
        auto_targeting = [r for r in rows if r["Entity"] == "Product Targeting"
                         and r["Product Targeting Expression Type"] == "auto"]
        manual_targeting = [r for r in rows if r["Entity"] == "Product Targeting"
                           and r["Product Targeting Expression Type"] == "manual"]
        assert len(auto_targeting) == 4
        assert len(manual_targeting) == 1

    def test_generate_with_no_skus_does_nothing(self, generator):
        generator.generate_clearance_campaigns("365+", [])
        assert len(generator.get_rows()) == 0

    def test_tier_budget_assignment(self, config):
        """365+ tier gets higher budget, 211-330 gets lower."""
        gen365 = BulkCampaignGenerator(config, DEFAULT_COLUMNS)
        gen365.generate_clearance_campaigns("365+", ["sku-001"])
        campaigns_365 = [r for r in gen365.get_rows() if r["Entity"] == "Campaign"]

        gen211 = BulkCampaignGenerator(config, DEFAULT_COLUMNS)
        gen211.generate_clearance_campaigns("211-330", ["sku-001"])
        campaigns_211 = [r for r in gen211.get_rows() if r["Entity"] == "Campaign"]

        assert campaigns_365[0]["Campaign Daily Budget"] == 50.0
        assert campaigns_211[0]["Campaign Daily Budget"] == 30.0


# =========================================================================
# XLSX Writer tests
# =========================================================================

class TestXLSXWriter:
    def test_write_and_read_back(self, generator):
        generator.add_campaign("Test Campaign", 50.0, "MANUAL")
        generator.add_bidding_adjustment(-1000, "placementTop", 80)
        rows = generator.get_rows()

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            output_path = f.name

        try:
            write_bulk_xlsx(rows, DEFAULT_COLUMNS, output_path)

            # Read back and verify
            from openpyxl import load_workbook
            wb = load_workbook(output_path)
            assert "Sponsored Products Campaigns" in wb.sheetnames
            ws = wb["Sponsored Products Campaigns"]

            # Check headers
            headers = [ws.cell(1, i).value for i in range(1, len(DEFAULT_COLUMNS) + 1)]
            assert headers == DEFAULT_COLUMNS

            # Check data rows
            assert ws.max_row == 3  # header + 2 data rows

            # Verify blank cells are None, not 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                for val in row:
                    assert val != 0, "Found a 0 value that should be blank"

            wb.close()
        finally:
            os.unlink(output_path)

    def test_creates_output_directory(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = os.path.join(tmpdir, "nested", "dir", "output.xlsx")
            write_bulk_xlsx([], DEFAULT_COLUMNS, output_path)
            assert os.path.exists(output_path)


# =========================================================================
# End-to-end tests
# =========================================================================

class TestEndToEnd:
    def test_full_generation(self, sample_listings_path, sample_tier_path):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            output_path = f.name

        with open(sample_tier_path) as f:
            tier_assignment = json.load(f)

        try:
            result = generate_bulk_upload(
                target_asins=["B08NV6CLGF", "B09ABC1234", "B09GHI9012"],
                listings_report_path=sample_listings_path,
                output_path=output_path,
                competitor_asins=["B07ABC1234"],
                tier_assignment=tier_assignment,
            )
            assert result.is_valid
            assert os.path.exists(output_path)

            # Verify file content
            from openpyxl import load_workbook
            wb = load_workbook(output_path)
            ws = wb["Sponsored Products Campaigns"]
            assert ws.max_row > 1  # At least headers + some data
            wb.close()
        finally:
            os.unlink(output_path)

    def test_missing_asin_handled_gracefully(self, sample_listings_path):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            output_path = f.name

        try:
            result = generate_bulk_upload(
                target_asins=["B08NV6CLGF", "NONEXISTENT1"],
                listings_report_path=sample_listings_path,
                output_path=output_path,
            )
            # Should still succeed with the valid ASIN
            assert result.is_valid
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_all_asins_missing_returns_error(self, sample_listings_path):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            output_path = f.name

        try:
            result = generate_bulk_upload(
                target_asins=["NONEXIST001", "NONEXIST002"],
                listings_report_path=sample_listings_path,
                output_path=output_path,
            )
            assert not result.is_valid
            assert any("No valid ASINs" in e for e in result.errors)
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_config_file_loading(self, sample_listings_path, sample_config_path):
        with open(sample_config_path) as f:
            config_dict = json.load(f)
        config = CampaignConfig(**config_dict)

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            output_path = f.name

        try:
            result = generate_bulk_upload(
                target_asins=["B08NV6CLGF"],
                listings_report_path=sample_listings_path,
                output_path=output_path,
                config=config,
            )
            assert result.is_valid
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)


# =========================================================================
# CampaignConfig tests
# =========================================================================

class TestCampaignConfig:
    def test_default_values(self):
        config = CampaignConfig()
        assert config.daily_budget_365_plus == 50.0
        assert config.daily_budget_211_330 == 30.0
        assert config.default_bid == 0.75
        assert config.bidding_strategy == "Dynamic bids - down only"
        assert config.placement_top_percentage == 80
        assert config.placement_product_page_percentage == 50

    def test_custom_values(self):
        config = CampaignConfig(
            daily_budget_365_plus=100.0,
            default_bid=1.50,
            campaign_prefix="Test",
        )
        assert config.daily_budget_365_plus == 100.0
        assert config.default_bid == 1.50
        assert config.campaign_prefix == "Test"

    def test_from_json(self, sample_config_path):
        with open(sample_config_path) as f:
            config_dict = json.load(f)
        config = CampaignConfig(**config_dict)
        assert config.daily_budget_365_plus == 50.0
        assert len(config.generic_keywords) == 7  # sample has 7 keywords
